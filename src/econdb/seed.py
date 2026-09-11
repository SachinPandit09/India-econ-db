"""`econdb seed`: load the two reference workbooks into the meta schema (idempotent upserts)."""

import re
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string
from psycopg.types.json import Jsonb

from econdb import db
from econdb.series_map import SOURCES, build, norm

REFERENCE = Path(__file__).resolve().parents[2] / "docs" / "reference"
LAYOUT = REFERENCE / "India_Econ_Tracker_Layout.xlsx"
CATALOGUE = REFERENCE / "India_Econ_Indicator_Catalogue.xlsx"
YELLOW = "FFFFF2CC"  # row-3 fill of computed change columns (docs/design.md §5)
FORMULA = re.compile(r'=IF\(AND\(ISNUMBER\(([A-Z]+)(\d+)\),ISNUMBER\(([A-Z]+)(\d+)\)\),(.+),""\)')
SHAPES = {"{a}/{b}-1": "pct", "({a}-{b})*100": "bps", "{a}-{b}": "diff"}
CATALOGUE_COLUMNS = {
    "ID": "catalogue_id",
    "Theme": "theme",
    "Indicator (family)": "family",
    "Key breakdowns": "breakdowns",
    "Frequency": "frequency",
    "Unit": "unit",
    "Source agency": "source_agency",
    "Access method": "access_method",
    "Access detail": "access_detail",
    "History": "history",
    "Verification": "verification",
    "Priority": "priority",
    "Build tier": "build_tier",
    "Notes": "notes",
    "Your decision": "decision",
    "Team notes": "team_notes",
}


@dataclass
class Column:
    sheet_code: str
    column_no: int
    block_title: str
    column_label: str
    calc_kind: str | None = None
    calc_lag: int | None = None
    calc_ref_column: int | None = None


def parse_formula(formula: str) -> tuple[str, int, int]:
    """'=IF(AND(ISNUMBER(C16),ISNUMBER(C4)),C16/C4-1,"")' -> ('pct', lag 12, column 3)."""
    m = FORMULA.fullmatch(formula)
    if m:
        col, row, ref_col, ref_row, body = m.groups()
        for shape, kind in SHAPES.items():
            if col == ref_col and body == shape.format(a=col + row, b=ref_col + ref_row):
                return kind, int(row) - int(ref_row), column_index_from_string(col)
    raise ValueError(f"unknown formula shape: {formula!r}")


def read_layout(path: Path = LAYOUT) -> tuple[list[dict], list[Column], list[dict]]:
    """CONTENTS rows, every non-period column of the 87 sheets, and O08 linking-factor pairs."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        sheets = [
            {
                "sheet_code": r[2][:3],
                "sheet_name": r[2],
                "position": r[0],
                "category": r[1],
                "title": r[3],
                "blocks": r[4],
                "coverage": r[5],
                "source_text": r[6],
                "access_text": r[7],
                "sheet_type": r[8],
                "derived_from": r[9],
                "status": r[10],
            }
            for r in wb["CONTENTS"].iter_rows(values_only=True)
            if isinstance(r[0], int)  # numbered sheet rows only
        ]
        columns = []
        for ws in wb.worksheets[2:]:
            code = ws.title[:3]
            rows = list(ws.iter_rows(max_row=60, max_col=ws.max_column))
            periods = 2 if code[0] in "QM" else 1  # Month + Fiscal year, Quarter + Months
            block = None
            for i, (top, head) in enumerate(zip(rows[1], rows[2], strict=True)):
                block = block if top.value is None else top.value
                if i < periods or head.value is None:
                    continue
                col = Column(code, i + 1, norm(block), norm(head.value))
                if head.fill.fill_type and head.fill.fgColor.rgb == YELLOW:
                    values = (r[i].value for r in rows[3:] if len(r) > i)
                    formula = next((v for v in values if isinstance(v, str) and v[:1] == "="), "")
                    col.calc_kind, col.calc_lag, col.calc_ref_column = parse_formula(formula)
                columns.append(col)
        linking = [
            {"family": r[1], "old_base": str(r[2]), "new_base": str(r[3])}
            for r in wb["O08_Linking_Factors"].iter_rows(min_row=4, values_only=True)
            if r[1]
        ]
    finally:
        wb.close()
    return sheets, columns, linking


def read_catalogue(path: Path = CATALOGUE) -> tuple[list[dict], dict[str, str], dict[int, tuple]]:
    """CATALOGUE rows, working Yahoo tickers (name -> ticker), RBI codes (code -> name/ccy/freq)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb["CATALOGUE"].iter_rows(values_only=True))
        catalogue = [
            {CATALOGUE_COLUMNS[h]: v for h, v in zip(rows[0], r, strict=True)}
            for r in rows[1:]
            if r[0]
        ]
        yahoo = wb["YAHOO_TICKERS"].iter_rows(values_only=True)
        tickers = {r[1]: r[0] for r in yahoo if r[3] == "Y"}
        rbi_rows = wb["RBI_API_39"].iter_rows(values_only=True)
        rbi = {r[0]: (r[1], r[2], r[3]) for r in rbi_rows if isinstance(r[0], int)}
    finally:
        wb.close()
    return catalogue, tickers, rbi


def series_rows(columns: list[Column], cells: dict, titles: dict[str, str]) -> list[dict]:
    """One row per distinct series; its name comes from the first column showing it natively."""
    found, native, fallback = {}, {}, {}
    for c in columns:
        cell = cells.get((c.sheet_code, c.block_title, c.column_label))
        if cell is None:
            continue
        s = cell.series
        if s.series_id in found and found[s.series_id] != s:
            raise RuntimeError(f"{s.series_id} is defined differently on {c.sheet_code}")
        found[s.series_id] = s
        name = f"{titles[c.sheet_code]} · {c.block_title} · {c.column_label}"
        names = native if cell.agg is None and cell.transform is None else fallback
        names.setdefault(s.series_id, name)
    return [
        {
            **asdict(s),
            "family": sid.split(".")[0],
            "freq": sid[-1].upper(),
            "name": native.get(sid) or fallback[sid],
        }
        for sid, s in found.items()
    ]


def upsert(conn, table: str, key: list[str], rows: list[dict], update: bool = True) -> int:
    """Insert new rows and update changed ones; returns the number of rows written."""
    cols = list(rows[0])
    other = [c for c in cols if c not in key]
    sql = (
        f"INSERT INTO {table} AS t ({', '.join(cols)}) VALUES ({', '.join(['%s'] * len(cols))})"
        f" ON CONFLICT ({', '.join(key)}) "
    )
    if update:
        sql += (
            f"DO UPDATE SET {', '.join(f'{c} = EXCLUDED.{c}' for c in other)}"
            f" WHERE ({', '.join(f't.{c}' for c in other)})"
            f" IS DISTINCT FROM ({', '.join(f'EXCLUDED.{c}' for c in other)})"
        )
    else:
        sql += "DO NOTHING"
    written = 0
    with conn.cursor() as cur:
        for row in rows:
            cur.execute(sql, [Jsonb(v) if isinstance(v, dict) else v for v in row.values()])
            written += cur.rowcount
    return written


def seed() -> int:
    sheets, columns, linking = read_layout()
    catalogue, tickers, rbi = read_catalogue()
    cells = build(tickers, rbi)
    stale = set(cells) - {(c.sheet_code, c.block_title, c.column_label) for c in columns}
    if stale:
        raise RuntimeError(f"series map entries not found in the layout: {sorted(stale)[:5]}")

    series = series_rows(columns, cells, {s["sheet_code"]: s["title"] for s in sheets})
    sheet_columns = []
    for c in columns:
        cell = cells.get((c.sheet_code, c.block_title, c.column_label))
        sheet_columns.append(
            {
                **asdict(c),
                "series_id": cell.series.series_id if cell else None,
                "agg": cell.agg if cell else None,
                "transform": cell.transform if cell else None,
            }
        )
    factors = [{**f, "scope": "all"} for f in linking]
    tables = {  # table: (key, rows, update existing rows?)
        "meta.source": (["source_id"], SOURCES, True),
        "meta.catalogue": (["catalogue_id"], catalogue, True),
        "meta.sheet": (["sheet_code"], sheets, True),
        "meta.series": (["series_id"], series, True),
        "meta.sheet_column": (["sheet_code", "column_no"], sheet_columns, True),
        # never overwrite a factor filled in later from an official notice
        "meta.linking_factor": (["family", "scope", "old_base", "new_base"], factors, False),
    }

    with db.connect() as conn:
        written = {
            t: upsert(conn, t, key, rows, update) for t, (key, rows, update) in tables.items()
        }
        removed = conn.execute(
            "DELETE FROM meta.sheet_column WHERE (sheet_code, column_no) NOT IN"
            " (SELECT * FROM unnest(%s::text[], %s::smallint[]))",
            ([c.sheet_code for c in columns], [c.column_no for c in columns]),
        ).rowcount

    for table, (_, rows, _) in tables.items():
        print(f"{table:<20} {len(rows):>5} rows  {written[table]:>5} written")
    if removed:
        print(f"meta.sheet_column    {removed:>5} stale rows removed")
    mapped = sum(1 for r in sheet_columns if r["series_id"])
    formulas = sum(1 for r in sheet_columns if r["calc_kind"])
    unmapped = len(sheet_columns) - mapped - formulas
    print(f"sheet columns: {mapped} mapped, {formulas} formulas, {unmapped} not mapped yet")
    for (source, dataset), n in sorted(
        Counter((s["source_id"], s["dataset"]) for s in series).items()
    ):
        print(f"series  {source:<6} {dataset:<9} {n:>4}")
    return 0
